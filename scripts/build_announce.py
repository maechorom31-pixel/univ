#!/usr/bin/env python3
"""
2027 수시 일정 달력(엑셀) → data/announce2027.json — 발표일과 대학별고사일.

왜 따로 만드나
--------------
1단계·최종 발표일은 여태 data/schedule2027.json(「수박먹고 대학간다」 7.18 판)에서
왔다. 그 표는 대학·전형 단위라 **학과가 갈리는 발표**를 못 담고, 전형 이름이 안 맞으면
같은 유형의 첫 줄을 집었다. 광주대 「간호학과만 12.16, 나머지는 10.29」 같은 것이
한 날짜로 뭉개졌다.

이 엑셀(「2027학년도 수시모집 일정 달력」 — 각 대학 2027 모집요강 전형일정 표에서 수집,
205개 대학)은 발표마다 **대상 전형·학과를 글로 적는다** — 「학생부교과(일반학생)
간호학과」, 「수능 미적용 학과」, 비고의 「간호학과 제외(12.16 발표)」. 그 글을 그대로
남기고, 고르는 일은 board/match.js 의 announceDates 가 한다.

전문대는 같은 파일의 「전형일정」 시트가 대학별 수시1차·2차 합격자 발표를 준다.

쓰는 법
-------
    python3 scripts/build_announce.py <일정 달력.xlsx> [<전문대 전형일정.xlsx>] \
        -o data/announce2027.json

두 시트가 한 파일에 있으면 그 파일 하나만 주면 된다. 원본 엑셀은 저장소에 넣지 않는다.
"""
import argparse
import datetime as dt
import json
import sys

try:
    import openpyxl
except ImportError:                                   # pragma: no cover
    sys.exit('openpyxl 이 필요합니다:  pip install openpyxl')

# 발표 둘과 대학별고사 셋. 고사는 「대표일」로 적히고 기간은 비고에 「~10.29」로 붙는다.
KIND = {'1단계합격': 'stage1', '최종합격': 'final', '면접': '면접', '논술': '논술', '실기': '실기'}


def day(v):
    """칸 값 → 'YYYY-MM-DD'. 날짜 칸·글 칸 둘 다 받는다. 못 읽으면 None."""
    if v is None or v == '':
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def text(v):
    return '' if v is None else str(v).strip()


def until(when, note):
    """비고의 「~10.29」「~10.29(목)」 → 끝나는 날. 없거나 앞날이면 시작일 그대로."""
    import re
    m = re.search(r'~\s*(\d{1,2})\.(\d{1,2})', note or '')
    if not m:
        return when
    end = f'{when[:4]}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'
    return end if end >= when else when


def read_list(ws):
    """일정목록 — 날짜·요일·구분·대학·전형·시각·비고. 발표 둘과 고사 셋만 남긴다."""
    rows = ws.iter_rows(values_only=True)
    head = [text(c) for c in next(rows)]
    want = ['날짜', '요일', '구분', '대학', '전형', '시각', '비고']
    if head[:7] != want:
        sys.exit(f'일정목록 머리가 예상과 다릅니다: {head[:7]}')
    out = []
    for r in rows:
        kind = KIND.get(text(r[2]))
        when = day(r[0])
        if not kind or not when:
            continue
        univ, note = text(r[3]), text(r[6])
        # 「국립목포대(미확정)」의 괄호는 캠퍼스가 아니라 상태다 — 캠퍼스로 읽으면
        # 어느 지원에도 안 맞는다. 이름에서 떼어 비고로 옮긴다.
        if univ.endswith('(미확정)'):
            univ = univ[:-len('(미확정)')]
            note = ('미확정 · ' + note).strip(' ·')
        out.append({
            'date': when, 'to': until(when, note), 'kind': kind, 'univ': univ,
            'scope': text(r[4]), 'clock': text(r[5]), 'note': note,
        })
    return out


def read_college(ws):
    """전형일정(전문대) — 대학명·모집시기·…·합격자 발표. 수시만 남긴다."""
    out = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        univ, period, final = text(r[3]), text(r[4]), day(r[10])
        if not univ or not period.startswith('수시') or not final:
            continue
        out.append({'univ': univ, 'period': period, 'final': final})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx', nargs='+')
    ap.add_argument('-o', '--out', required=True)
    a = ap.parse_args()

    rows, college = None, None
    for path in a.xlsx:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        if '일정목록' in names:
            got = read_list(wb['일정목록'])
            if rows is not None and got != rows:
                sys.exit('두 파일의 일정목록이 다릅니다 — 어느 쪽이 최신인지 먼저 정해 주세요.')
            rows = got
        if '전형일정' in names:
            college = read_college(wb['전형일정'])
    if rows is None:
        sys.exit('일정목록 시트를 찾지 못했습니다.')

    doc = {
        'year': 2027,
        'source': '2027학년도 수시모집 일정 달력(각 대학 2027 수시모집요강 전형일정 표, 2026년 6월 수집)'
                  ' · 전문대 전형일정(수시1·2차 합격자 발표)',
        'rows': sorted(rows, key=lambda r: (r['date'], r['univ'], r['kind'])),
        'college': college or [],
    }
    with open(a.out, 'w', encoding='utf-8') as f:
        json.dump(doc, f, ensure_ascii=False, separators=(',', ':'))
    n = lambda k: sum(r['kind'] == k for r in rows)
    print(f'{len(rows)}줄 — 1단계 {n("stage1")} · 최종 {n("final")} · 면접 {n("면접")}'
          f' · 논술 {n("논술")} · 실기 {n("실기")} · 전문대 {len(doc["college"])}줄 → {a.out}')


if __name__ == '__main__':
    main()
